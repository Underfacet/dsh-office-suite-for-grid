#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""P2 Office Python 编辑/排版引擎（python-docx / openpyxl）

调用：<venv python> office-python.py -Json '{"mode":"edit"|"format", ...}'
      也支持从 stdin 读取同一 JSON（未提供 -Json 时）。

输入 payload：
  edit:   {mode:'edit', path, ext, app, ops:[...], outputPath, create}
  format: {mode:'format', path, ext, app:'word', spec:{rules,scope}, templatePath, placeholders, outputPath}

输出（stdout 仅一行 JSON）：
  {ok:true,result:{outputPath,saved,applied,stats,warnings?}}
  {ok:false,error:{code,message,hint?},result?:{...}}
"""

import json
import os
import re
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass


class OpError(Exception):
    def __init__(self, code, message, hint=None):
        super().__init__(message)
        self.code = code
        self.message = message
        self.hint = hint


def emit(payload, exit_code=0):
    sys.stdout.write(json.dumps(payload, ensure_ascii=False))
    sys.stdout.flush()
    raise SystemExit(exit_code)


def read_payload(argv):
    if "-Json" in argv:
        index = argv.index("-Json")
        if index + 1 >= len(argv):
            raise OpError("PARAM_INVALID", "-Json 缺少参数值")
        return json.loads(argv[index + 1])
    raw = sys.stdin.buffer.read().decode("utf-8-sig")
    return json.loads(raw) if raw.strip() else {}


# ───────────────────────── 通用工具 ─────────────────────────

def chinese_number(n):
    digits = "零一二三四五六七八九"
    if n <= 0:
        return "零"
    if n < 10:
        return digits[n]
    if n == 10:
        return "十"
    if n < 20:
        return "十" + digits[n - 10]
    if n < 100:
        tens, units = divmod(n, 10)
        return digits[tens] + "十" + (digits[units] if units else "")
    return str(n)


NUMBER_PREFIX_RE = re.compile(
    r"^\s*(?:[一二三四五六七八九十百零〇]+[、.．]|（[一二三四五六七八九十百零〇]+）|\([0-9一二三四五六七八九十]+\)|[0-9]+(?:\.[0-9]+)*[、.．]?\s*)"
)


def strip_number_prefix(text):
    return NUMBER_PREFIX_RE.sub("", text or "", count=1)


def number_prefix(level, counters, mode):
    if mode in (None, "", "none"):
        return ""
    if mode == "chinese":
        cn = chinese_number(counters.get(level, 0))
        if level == 1:
            return cn + "、"
        if level == 2:
            return "（" + cn + "）"
        return str(counters.get(level, 0)) + ". "
    if mode == "arabic":
        return str(counters.get(level, 0)) + ". "
    if mode == "decimalDot":
        return ".".join(str(counters.get(i, 0)) for i in range(1, level + 1)) + " "
    return ""


def heading_level(paragraph):
    name = ""
    try:
        name = paragraph.style.name or ""
    except Exception:
        name = ""
    match = re.match(r"^(?:标题|Heading)\s*([1-9])$", name)
    if match:
        return int(match.group(1))
    try:
        outline = paragraph.paragraph_format.outline_level
        if outline is not None and 1 <= int(outline) <= 9:
            return int(outline)
    except Exception:
        pass
    return 0


def resolve_style(doc, name):
    try:
        return doc.styles[name]
    except Exception:
        pass
    match = re.match(r"^(?:标题|Heading)\s*([1-9])$", name or "")
    if match:
        level = match.group(1)
        for candidate in ("Heading " + level, "标题 " + level):
            try:
                return doc.styles[candidate]
            except Exception:
                pass
    try:
        for style in doc.styles:
            if style.name == name:
                return style
    except Exception:
        pass
    if name in ("正文", "Normal"):
        for candidate in ("Normal", "正文"):
            try:
                return doc.styles[candidate]
            except Exception:
                pass
    raise OpError("E2005", "样式不存在: %s" % name)


def apply_font_to_paragraph(paragraph, font):
    if not font:
        return
    from docx.shared import Pt
    from docx.oxml.ns import qn

    for run in paragraph.runs:
        if font.get("name"):
            run.font.name = font["name"]
            try:
                rpr = run._element.get_or_add_rPr()
                rfonts = rpr.find(qn("w:rFonts"))
                if rfonts is None:
                    rfonts = rpr.makeelement(qn("w:rFonts"), {})
                    rpr.append(rfonts)
                rfonts.set(qn("w:eastAsia"), font["name"])
                rfonts.set(qn("w:ascii"), font["name"])
                rfonts.set(qn("w:hAnsi"), font["name"])
            except Exception:
                pass
        if font.get("sizePt") is not None:
            run.font.size = Pt(float(font["sizePt"]))
        if font.get("bold") is not None:
            run.font.bold = bool(font["bold"])


def apply_paragraph_format(paragraph, fmt):
    if not fmt:
        return
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    pf = paragraph.paragraph_format
    if fmt.get("lineSpacing") is not None:
        pf.line_spacing = float(fmt["lineSpacing"])
    elif fmt.get("lineSpacingPt") is not None:
        pf.line_spacing = Pt(float(fmt["lineSpacingPt"]))
    if fmt.get("spaceBeforePt") is not None:
        pf.space_before = Pt(float(fmt["spaceBeforePt"]))
    if fmt.get("spaceAfterPt") is not None:
        pf.space_after = Pt(float(fmt["spaceAfterPt"]))
    aligns = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    if fmt.get("align") in aligns:
        pf.alignment = aligns[fmt["align"]]


def word_stats(doc):
    tables = len(doc.tables)
    sections = len(doc.sections)
    paragraphs = len(doc.paragraphs)
    words = 0
    for paragraph in doc.paragraphs:
        words += len((paragraph.text or "").split())
    return {"paragraphs": paragraphs, "tables": tables, "sections": sections, "words": words}


# ───────────────────────── Word 编辑 ─────────────────────────

def word_edit(payload):
    from docx import Document

    path = payload["path"]
    ext = (payload.get("ext") or os.path.splitext(path)[1]).lower()
    if ext == ".doc":
        raise OpError("E2004", "python-docx 不支持 .doc 旧格式", "请安装 Word 使用 engine=com，或先另存为 .docx")
    if os.path.exists(path) and payload.get("create"):
        raise OpError("E2005", "create 仅用于不存在的文件: %s" % path)
    if payload.get("create"):
        doc = Document()
    else:
        try:
            doc = Document(path)
        except Exception as err:  # noqa: BLE001
            raise OpError("E2004", "打开 Word 文档失败: %s" % err, "请确认文件未损坏/未被 Office 占用")

    applied = []
    failure = None
    ops = payload.get("ops") or []
    for index, op in enumerate(ops):
        action = op.get("action")
        try:
            invoke_word_op(doc, op)
            applied.append({"op": action, "index": index, "ok": True})
        except OpError as err:
            applied.append({"op": action, "index": index, "ok": False, "message": err.message})
            failure = err
            break
        except Exception as err:  # noqa: BLE001 - 引擎统一映射
            applied.append({"op": action, "index": index, "ok": False, "message": str(err)})
            failure = OpError("E2005", "执行 %s 失败: %s" % (action, err))
            break

    output_path = payload.get("outputPath") or path
    saved = False
    try:
        ensure_parent(output_path)
        doc.save(output_path)
        saved = True
    except Exception as err:  # noqa: BLE001
        if failure is None:
            failure = OpError("E2010", "保存文档失败: %s" % err)
        saved = False

    stats = word_stats(doc)
    if failure is not None:
        return {
            "ok": False,
            "error": {"code": failure.code, "message": failure.message, **({"hint": failure.hint} if failure.hint else {})},
            "result": {"outputPath": output_path, "saved": saved, "applied": applied, "stats": stats},
        }
    return {"ok": True, "result": {"outputPath": output_path, "saved": True, "applied": applied, "stats": stats}}


def invoke_word_op(doc, op):
    action = op.get("action")
    if action in ("create", "saveAs"):
        return
    if action == "paragraph":
        sub = op.get("op")
        if sub == "append":
            paragraph = doc.add_paragraph(op.get("text") or "")
            return
        at = int(op.get("at", 0))
        count = len(doc.paragraphs)
        if at < 1 or at > count:
            raise OpError("E2005", "段落索引 %d 超出范围（共 %d 段）" % (at, count))
        paragraph = doc.paragraphs[at - 1]
        if sub == "setText":
            paragraph.text = op.get("text") or ""
        elif sub == "setStyle":
            paragraph.style = resolve_style(doc, op.get("style"))
        elif sub == "insert":
            new_para = paragraph.insert_paragraph_before(op.get("text") or "")
            if op.get("style"):
                new_para.style = resolve_style(doc, op.get("style"))
        elif sub == "delete":
            element = paragraph._element
            element.getparent().remove(element)
        else:
            raise OpError("E2005", "未知段落操作: %s" % sub)
        return
    if action == "heading":
        sub = op.get("op")
        level = int(op.get("level", 0))
        if level < 1 or level > 3:
            raise OpError("E2005", "标题级别必须为 1–3: %s" % level)
        style = resolve_style(doc, "Heading %d" % level)
        at = int(op.get("at", 0))
        count = len(doc.paragraphs)
        if at < 1 or at > count:
            raise OpError("E2005", "段落索引 %d 超出范围（共 %d 段）" % (at, count))
        paragraph = doc.paragraphs[at - 1]
        if sub == "setText":
            paragraph.text = op.get("text") or ""
            paragraph.style = style
        elif sub == "setStyle":
            paragraph.style = resolve_style(doc, op.get("style")) if op.get("style") else style
        elif sub == "insert":
            new_para = paragraph.insert_paragraph_before(op.get("text") or "")
            new_para.style = style
        elif sub == "delete":
            element = paragraph._element
            element.getparent().remove(element)
        else:
            raise OpError("E2005", "未知标题操作: %s" % sub)
        return
    if action == "table":
        invoke_word_table_op(doc, op)
        return
    raise OpError("E2005", "未知 Word 指令: %s" % action)


def get_table(doc, index):
    tables = doc.tables
    if index < 1 or index > len(tables):
        raise OpError("E2005", "表格索引 %d 超出范围（共 %d 个表格）" % (index, len(tables)))
    return tables[index - 1]


def invoke_word_table_op(doc, op):
    table = get_table(doc, int(op.get("table", 0)))
    sub = op.get("op")
    if sub == "setCell":
        row = int(op.get("row", 0))
        col = int(op.get("col", 0))
        if row < 1 or row > len(table.rows):
            raise OpError("E2005", "行索引 %d 超出范围（共 %d 行）" % (row, len(table.rows)))
        if col < 1 or col > len(table.columns):
            raise OpError("E2005", "列索引 %d 超出范围（共 %d 列）" % (col, len(table.columns)))
        table.cell(row - 1, col - 1).text = op.get("text") or ""
        return
    if sub == "addRow":
        if op.get("at") not in (None, "", "end"):
            at = int(op.get("at"))
            if at < 1 or at > len(table.rows):
                raise OpError("E2005", "插入行位置 %d 超出范围" % at)
            target = table.rows[at - 1]._tr
            new_row = table.add_row()
            target.addprevious(new_row._tr)
        else:
            table.add_row()
        return
    if sub == "deleteRow":
        row = int(op.get("row", 0))
        if row < 1 or row > len(table.rows):
            raise OpError("E2005", "行索引 %d 超出范围" % row)
        table._tbl.remove(table.rows[row - 1]._tr)
        return
    if sub == "addColumn":
        from docx.shared import Mm

        table.add_column(Mm(20))
        return
    if sub == "deleteColumn":
        col = int(op.get("col", 0))
        grid = table._tbl.tblGrid
        from docx.oxml.ns import qn

        grid_cols = grid.findall(qn("w:gridCol"))
        if col < 1 or col > len(grid_cols):
            raise OpError("E2005", "列索引 %d 超出范围（共 %d 列）" % (col, len(grid_cols)))
        grid.remove(grid_cols[col - 1])
        for row in table.rows:
            cells = row._tr.findall(qn("w:tc"))
            if col - 1 < len(cells):
                row._tr.remove(cells[col - 1])
        return
    raise OpError("E2005", "未知表格操作: %s" % sub)


# ───────────────────────── Word 排版 ─────────────────────────

def word_format(payload):
    from docx import Document

    path = payload["path"]
    ext = (payload.get("ext") or os.path.splitext(path)[1]).lower()
    if ext == ".doc":
        raise OpError("E2004", "python-docx 不支持 .doc 旧格式", "请安装 Word 使用 engine=com，或先另存为 .docx")

    output_path = payload.get("outputPath") or path
    template_path = payload.get("templatePath") or ""
    placeholders = payload.get("placeholders") or {}

    if template_path:
        try:
            doc = Document(template_path)
        except Exception as err:  # noqa: BLE001
            raise OpError("E2007", "打开模板失败: %s" % err, "请确认模板 .docx 未损坏")
        replaced = expand_placeholders(doc, placeholders)
        ensure_parent(output_path)
        doc.save(output_path)
        stats = word_stats(doc)
        return {
            "ok": True,
            "result": {
                "outputPath": output_path,
                "saved": True,
                "applied": [{"rule": "template", "target": "template", "count": replaced, "ok": True}],
                "stats": stats,
                "templatePlaceholders": replaced,
            },
        }

    try:
        doc = Document(path)
    except Exception as err:  # noqa: BLE001
        raise OpError("E2004", "打开 Word 文档失败: %s" % err, "请确认文件未损坏/未被 Office 占用")
    spec = payload.get("spec") or {}
    rules = spec.get("rules") or []
    scope = spec.get("scope") or {"type": "all"}

    apply_word_numbering(doc, rules, scope)

    applied = []
    failure = None
    for index, rule in enumerate(rules):
        target = rule.get("target")
        try:
            count = apply_word_rule(doc, rule, scope)
            entry = {"rule": index, "target": target, "count": count, "ok": True}
            if target == "heading":
                entry["level"] = int(rule.get("level", 0))
            applied.append(entry)
        except OpError as err:
            applied.append({"rule": index, "target": target, "count": 0, "ok": False, "message": err.message})
            failure = err
            break
        except Exception as err:  # noqa: BLE001
            applied.append({"rule": index, "target": target, "count": 0, "ok": False, "message": str(err)})
            failure = OpError("E2006", "应用排版规则失败: %s" % err)
            break

    saved = False
    try:
        ensure_parent(output_path)
        doc.save(output_path)
        saved = True
    except Exception as err:  # noqa: BLE001
        if failure is None:
            failure = OpError("E2010", "保存文档失败: %s" % err)
        saved = False

    stats = word_stats(doc)
    if failure is not None:
        return {
            "ok": False,
            "error": {"code": failure.code, "message": failure.message, **({"hint": failure.hint} if failure.hint else {})},
            "result": {"outputPath": output_path, "saved": saved, "applied": applied, "stats": stats},
        }
    return {"ok": True, "result": {"outputPath": output_path, "saved": True, "applied": applied, "stats": stats}}


def apply_word_numbering(doc, rules, scope):
    numbering_by_level = {}
    for rule in rules:
        if rule.get("target") != "heading":
            continue
        if "numbering" not in rule:
            continue
        level = int(rule.get("level", 0))
        if 1 <= level <= 3:
            numbering_by_level[level] = rule.get("numbering")
    if not numbering_by_level:
        return
    if scope.get("type") == "paragraphs":
        return

    counters = {1: 0, 2: 0, 3: 0}
    allowed = scope.get("headings") if scope.get("type") == "headings" else None
    for paragraph in doc.paragraphs:
        level = heading_level(paragraph)
        if level < 1 or level > 3:
            continue
        if allowed is not None and level not in allowed:
            continue
        counters[level] = counters.get(level, 0) + 1
        for deeper in range(level + 1, 4):
            counters[deeper] = 0
        if level not in numbering_by_level:
            continue
        mode = numbering_by_level[level]
        if mode in (None, "none"):
            continue
        prefix = number_prefix(level, counters, mode)
        paragraph.text = prefix + strip_number_prefix(paragraph.text)


def apply_word_rule(doc, rule, scope):
    target = rule.get("target")
    if target == "page":
        return apply_page_rule(doc, rule)
    if target in ("header", "footer"):
        return apply_header_footer_rule(doc, rule)

    count = 0
    if scope.get("type") == "paragraphs":
        paragraphs = doc.paragraphs
        for idx in scope.get("paragraphs") or []:
            if idx < 1 or idx > len(paragraphs):
                raise OpError("E2006", "scope.paragraphs 索引 %d 超出范围（共 %d 段）" % (idx, len(paragraphs)))
            paragraph = paragraphs[idx - 1]
            apply_font_to_paragraph(paragraph, rule.get("font"))
            apply_paragraph_format(paragraph, rule.get("paragraph"))
            count += 1
        return count

    if target == "heading":
        level = int(rule.get("level", 0))
        allowed = scope.get("headings") if scope.get("type") == "headings" else None
        if allowed is not None and level not in allowed:
            return 0
        for paragraph in doc.paragraphs:
            if heading_level(paragraph) != level:
                continue
            apply_font_to_paragraph(paragraph, rule.get("font"))
            apply_paragraph_format(paragraph, rule.get("paragraph"))
            count += 1
        return count

    if target == "body":
        for paragraph in doc.paragraphs:
            if heading_level(paragraph) != 0:
                continue
            if not (paragraph.text or "").strip():
                continue
            apply_font_to_paragraph(paragraph, rule.get("font"))
            apply_paragraph_format(paragraph, rule.get("paragraph"))
            count += 1
        return count

    raise OpError("E2006", "未知排版目标: %s" % target)


def apply_page_rule(doc, rule):
    from docx.shared import Mm

    page = rule.get("page") or {}
    margins = page.get("marginsMm") or {}
    count = 0
    for section in doc.sections:
        if margins.get("top") is not None:
            section.top_margin = Mm(float(margins["top"]))
        if margins.get("bottom") is not None:
            section.bottom_margin = Mm(float(margins["bottom"]))
        if margins.get("left") is not None:
            section.left_margin = Mm(float(margins["left"]))
        if margins.get("right") is not None:
            section.right_margin = Mm(float(margins["right"]))
        count += 1
    return count


def apply_header_footer_rule(doc, rule):
    is_header = rule.get("target") == "header"
    payload = rule.get("header") if is_header else rule.get("footer")
    if not payload:
        return 0
    text = payload.get("text") or ""
    count = 0
    for section in doc.sections:
        container = section.header if is_header else section.footer
        paragraphs = container.paragraphs
        if not paragraphs:
            paragraph = container.add_paragraph()
        else:
            paragraph = paragraphs[0]
        paragraph.text = text
        if payload.get("font"):
            apply_font_to_paragraph(paragraph, payload.get("font"))
        count += 1
    return count


def expand_placeholders(doc, placeholders):
    if not placeholders:
        return 0
    count = 0

    def replace_paragraph(paragraph):
        nonlocal count
        for key, value in placeholders.items():
            for token in ("【%s】" % key, "{{%s}}" % key):
                if token not in (paragraph.text or ""):
                    continue
                replaced = False
                for run in paragraph.runs:
                    if token in run.text:
                        run.text = run.text.replace(token, str(value))
                        replaced = True
                if not replaced:
                    paragraph.text = (paragraph.text or "").replace(token, str(value))
                count += 1

    for paragraph in doc.paragraphs:
        replace_paragraph(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_paragraph(paragraph)
    for section in doc.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                replace_paragraph(paragraph)
    return count


# ───────────────────────── Excel ─────────────────────────

def excel_edit(payload):
    import openpyxl
    from openpyxl.utils import column_index_from_string

    path = payload["path"]
    ext = (payload.get("ext") or os.path.splitext(path)[1]).lower()
    if ext == ".xls":
        raise OpError("E2004", "openpyxl 不支持 .xls 旧格式", "请安装 Excel 使用 engine=com，或先另存为 .xlsx")
    if os.path.exists(path) and payload.get("create"):
        raise OpError("E2005", "create 仅用于不存在的文件: %s" % path)
    if payload.get("create"):
        workbook = openpyxl.Workbook()
    else:
        try:
            workbook = openpyxl.load_workbook(path)
        except Exception as err:  # noqa: BLE001
            raise OpError("E2004", "打开 Excel 工作簿失败: %s" % err, "请确认文件未损坏/未被 Office 占用")

    applied = []
    failure = None
    for index, op in enumerate(payload.get("ops") or []):
        action = op.get("action")
        try:
            invoke_excel_op(workbook, op, column_index_from_string)
            applied.append({"op": action, "index": index, "ok": True})
        except OpError as err:
            applied.append({"op": action, "index": index, "ok": False, "message": err.message})
            failure = err
            break
        except Exception as err:  # noqa: BLE001
            applied.append({"op": action, "index": index, "ok": False, "message": str(err)})
            failure = OpError("E2005", "执行 %s 失败: %s" % (action, err))
            break

    output_path = payload.get("outputPath") or path
    saved = False
    try:
        ensure_parent(output_path)
        workbook.save(output_path)
        saved = True
    except Exception as err:  # noqa: BLE001
        if failure is None:
            failure = OpError("E2010", "保存工作簿失败: %s" % err)
        saved = False

    stats = excel_stats(workbook)
    if failure is not None:
        return {
            "ok": False,
            "error": {"code": failure.code, "message": failure.message, **({"hint": failure.hint} if failure.hint else {})},
            "result": {"outputPath": output_path, "saved": saved, "applied": applied, "stats": stats},
        }
    return {"ok": True, "result": {"outputPath": output_path, "saved": True, "applied": applied, "stats": stats}}


def get_worksheet(workbook, name):
    if not name:
        return workbook.active
    if name not in workbook.sheetnames:
        raise OpError("E2005", "工作表不存在: %s" % name)
    return workbook[name]


def invoke_excel_op(workbook, op, column_index_from_string):
    action = op.get("action")
    if action in ("create", "saveAs"):
        return
    if action == "cell":
        worksheet = get_worksheet(workbook, op.get("sheet"))
        cell = worksheet[op.get("at")]
        if op.get("formula"):
            cell.value = op["formula"]
        else:
            cell.value = op.get("value")
        return
    if action == "row":
        worksheet = get_worksheet(workbook, op.get("sheet"))
        at = int(op.get("at", 0))
        if at < 1:
            raise OpError("E2005", "行号必须为正整数: %s" % at)
        if op.get("op") == "insert":
            worksheet.insert_rows(at, 1)
        elif op.get("op") == "delete":
            worksheet.delete_rows(at, 1)
        else:
            raise OpError("E2005", "未知行操作: %s" % op.get("op"))
        return
    if action == "column":
        worksheet = get_worksheet(workbook, op.get("sheet"))
        raw = op.get("at")
        if isinstance(raw, str) and re.match(r"^[A-Za-z]{1,3}$", raw):
            index = column_index_from_string(raw.upper())
        else:
            index = int(raw)
        if index < 1:
            raise OpError("E2005", "列号必须为正整数: %s" % index)
        if op.get("op") == "insert":
            worksheet.insert_cols(index, 1)
        elif op.get("op") == "delete":
            worksheet.delete_cols(index, 1)
        else:
            raise OpError("E2005", "未知列操作: %s" % op.get("op"))
        return
    if action == "sheet":
        sub = op.get("op")
        name = op.get("name")
        if sub == "rename":
            if name not in workbook.sheetnames:
                raise OpError("E2005", "工作表不存在: %s" % name)
            workbook[name].title = op.get("newName")
        elif sub == "add":
            workbook.create_sheet(name)
        elif sub == "delete":
            if name not in workbook.sheetnames:
                raise OpError("E2005", "工作表不存在: %s" % name)
            if len(workbook.sheetnames) <= 1:
                raise OpError("E2005", "不能删除工作簿中的唯一工作表")
            del workbook[name]
        else:
            raise OpError("E2005", "未知工作表操作: %s" % sub)
        return
    raise OpError("E2005", "未知 Excel 指令: %s" % action)


def excel_stats(workbook):
    try:
        sheet = workbook.active
        return {
            "sheets": len(workbook.sheetnames),
            "rows": sheet.max_row,
            "cols": sheet.max_column,
        }
    except Exception:
        return {"sheets": len(workbook.sheetnames)}


# ───────────────────────── 入口 ─────────────────────────

def ensure_parent(path):
    parent = os.path.dirname(os.path.abspath(path))
    if parent and not os.path.isdir(parent):
        os.makedirs(parent, exist_ok=True)


def main():
    try:
        payload = read_payload(sys.argv[1:])
        mode = payload.get("mode")
        path = payload.get("path") or ""
        ext = (payload.get("ext") or os.path.splitext(path)[1]).lower()
        if not path:
            raise OpError("E1041", "缺少 path")
        if mode not in ("edit", "format"):
            raise OpError("E1041", "mode 必须为 'edit' | 'format'")

        if ext in (".doc", ".docx"):
            result = word_edit(payload) if mode == "edit" else word_format(payload)
        elif ext in (".xls", ".xlsx"):
            if mode == "format":
                raise OpError("E2006", "Excel 排版不在 P2 范围（仅 Word）")
            result = excel_edit(payload)
        else:
            raise OpError("E1042", "不支持的 Office 文件类型: %s" % (ext or "(无扩展名)"))
        emit(result, 0 if result.get("ok") else 1)
    except OpError as err:
        emit({"ok": False, "error": {"code": err.code, "message": err.message, **({"hint": err.hint} if err.hint else {})}}, 1)
    except SystemExit:
        raise
    except Exception as err:  # noqa: BLE001
        import traceback

        emit(
            {
                "ok": False,
                "error": {
                    "code": "E1051",
                    "message": "Python Office 引擎异常: %s" % err,
                    "hint": traceback.format_exc(limit=3)[-800:],
                },
            },
            1,
        )


if __name__ == "__main__":
    main()

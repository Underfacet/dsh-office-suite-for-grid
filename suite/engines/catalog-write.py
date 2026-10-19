#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""P3 图纸目录表写出引擎（openpyxl；缺失时降级 CSV）

调用（三种取参方式，优先级从高到低）：
  1) <venv python> catalog-write.py -JsonFile '<payload.json 路径>'   ← Host 默认（避开 Windows 32K 命令行上限）
  2) <venv python> catalog-write.py -Json '{"path":...}'
  3) 从 stdin 读取同一 JSON（`type payload.json | python catalog-write.py`）

输入 payload：
  {
    "path": "输出 .xlsx 绝对路径",                 // 必填
    "title": "图纸目录表",                          // 可选（保留给表头语境）
    "headers": ["序号","图号","图名","图幅","比例","张数","来源文件"],  // 可选，缺省即该 7 列
    "sheets": [
      {"name": "汇总", "rows": [
        {"seq":1,"drawingNo":"D0101-01","title":"主接线","frameSize":"A4",
         "frameSizeLabel":"A4（横）","orientation":"landscape","scale":"1:100",
         "count":1,"source":"盘溪变-二次图纸.pdf"}
      ]}
    ]
  }

输出（stdout 仅一行 JSON）：
  {ok:true,result:{outputPath,engine:'python'|'csv',sheets:[...],rows:N,warnings?}}
  {ok:false,error:{code,message,hint?}}

降级：openpyxl 缺失/导入失败 → 写同目录同名 .csv（UTF-8 带 BOM，仅「汇总」表），
      结果 engine='csv' 并附 warnings（无样式、单工作表）。
"""

import csv
import json
import os
import re
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

DEFAULT_HEADERS = ["序号", "图号", "图名", "图幅", "比例", "张数", "来源文件"]

# 表头 → 行字段（帧幅列优先用 frameSizeLabel，回落 frameSize）
HEADER_FIELDS = {
    "序号": "seq",
    "图号": "drawingNo",
    "图名": "title",
    "图幅": "frameSizeLabel",
    "比例": "scale",
    "张数": "count",
    "来源文件": "source",
}

# 列宽（其余列 10）
COLUMN_WIDTHS = {"序号": 6, "图号": 18, "图名": 32}
DEFAULT_WIDTH = 10

# Excel 不允许的控制字符（openpyxl 会抛 IllegalCharacterError）
ILLEGAL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class OpError(Exception):
    def __init__(self, code, message, hint=None):
        super().__init__(message)
        self.code = code
        self.message = message
        self.hint = hint


def emit(payload, exit_code=0):
    sys.stdout.write(json.dumps(payload, ensure_ascii=False))
    sys.stdout.flush()
    raise SystemExit(exit_code)


def read_payload(argv):
    if "-JsonFile" in argv:
        index = argv.index("-JsonFile")
        if index + 1 >= len(argv):
            raise OpError("E1041", "-JsonFile 缺少参数值")
        with open(argv[index + 1], "r", encoding="utf-8-sig") as handle:
            return json.load(handle)
    if "-Json" in argv:
        index = argv.index("-Json")
        if index + 1 >= len(argv):
            raise OpError("E1041", "-Json 缺少参数值")
        return json.loads(argv[index + 1])
    raw = sys.stdin.buffer.read().decode("utf-8-sig")
    return json.loads(raw) if raw.strip() else {}


def sanitize(value):
    """单元格值规整：去控制字符；None 保持空。"""
    if value is None:
        return None
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (int, float)):
        return value
    return ILLEGAL_RE.sub("", str(value))


def normalize(payload):
    if not isinstance(payload, dict):
        raise OpError("E1041", "payload 必须是 JSON 对象")
    path = payload.get("path")
    if not isinstance(path, str) or not path.strip():
        raise OpError("E1041", "缺少 path（输出 .xlsx 路径）")
    path = os.path.abspath(path.strip())
    headers = payload.get("headers")
    if not isinstance(headers, list) or not headers:
        headers = list(DEFAULT_HEADERS)
    headers = [str(h) for h in headers]
    sheets = payload.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise OpError("E1041", "缺少 sheets（至少一个工作表）")
    normalized = []
    used = set()
    for index, sheet in enumerate(sheets):
        if not isinstance(sheet, dict):
            raise OpError("E1041", "sheets[%d] 必须是对象" % index)
        name = str(sheet.get("name") or "Sheet%d" % (index + 1)).strip() or "Sheet%d" % (index + 1)
        # Excel 工作表名限制：≤31 字符、不含 []:*?/\
        name = re.sub(r"[\[\]:*?/\\]", "_", name)[:31]
        base = name
        suffix = 1
        while name in used:
            suffix += 1
            tail = "(%d)" % suffix
            name = (base[: 31 - len(tail)] + tail) if len(base) + len(tail) > 31 else base + tail
        used.add(name)
        rows = sheet.get("rows")
        normalized.append({"name": name, "rows": rows if isinstance(rows, list) else []})
    return {"path": path, "title": payload.get("title"), "headers": headers, "sheets": normalized}


def cell_value(row, header):
    field = HEADER_FIELDS.get(header, header)
    value = row.get(field)
    if value is None and field == "frameSizeLabel":
        value = row.get("frameSize")
    if value is None and header == "序号":
        value = row.get("seq")
    return sanitize(value)


def write_xlsx(payload):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True)
    total_rows = 0
    sheet_names = []

    for sheet in payload["sheets"]:
        worksheet = workbook.create_sheet(title=sheet["name"])
        sheet_names.append(sheet["name"])
        headers = payload["headers"]
        worksheet.append(headers)
        for column, header in enumerate(headers, start=1):
            cell = worksheet.cell(row=1, column=column)
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            width = COLUMN_WIDTHS.get(header, DEFAULT_WIDTH)
            worksheet.column_dimensions[get_column_letter(column)].width = width
        for index, row in enumerate(sheet["rows"], start=1):
            if not isinstance(row, dict):
                continue
            values = [cell_value(row, header) for header in headers]
            worksheet.append(values)
            excel_row = index + 1
            for column in range(1, len(headers) + 1):
                cell = worksheet.cell(row=excel_row, column=column)
                cell.border = border
                if column == 1:
                    cell.alignment = Alignment(horizontal="center")
            total_rows += 1
        worksheet.freeze_panes = "A2"

    target = payload["path"]
    directory = os.path.dirname(target)
    if directory:
        os.makedirs(directory, exist_ok=True)
    workbook.save(target)
    return {
        "outputPath": target,
        "engine": "python",
        "sheets": sheet_names,
        "rows": total_rows,
    }


def write_csv(payload):
    target = os.path.splitext(payload["path"])[0] + ".csv"
    directory = os.path.dirname(target)
    if directory:
        os.makedirs(directory, exist_ok=True)
    headers = payload["headers"]
    total_rows = 0
    primary = payload["sheets"][0]
    with open(target, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        for row in primary["rows"]:
            if not isinstance(row, dict):
                continue
            writer.writerow(["" if cell_value(row, header) is None else cell_value(row, header) for header in headers])
            total_rows += 1
    return {
        "outputPath": target,
        "engine": "csv",
        "sheets": [primary["name"]],
        "rows": total_rows,
        "warnings": ["openpyxl 不可用，已降级 CSV（UTF-8 BOM）：无表头样式/冻结行，且仅输出「%s」工作表" % primary["name"]],
    }


def main(argv):
    try:
        payload = normalize(read_payload(argv))
    except OpError as err:
        emit({"ok": False, "error": {"code": err.code, "message": err.message, **({"hint": err.hint} if err.hint else {})}}, 0)
    except json.JSONDecodeError as err:
        emit({"ok": False, "error": {"code": "E1041", "message": "payload 不是合法 JSON: %s" % err}}, 0)
    except OSError as err:
        emit({"ok": False, "error": {"code": "E1041", "message": "读取 payload 失败: %s" % err}}, 0)

    try:
        result = write_xlsx(payload)
    except ImportError:
        try:
            result = write_csv(payload)
        except OSError as err:
            emit({"ok": False, "error": {"code": "E1062", "message": "目录表写盘失败: %s" % err, "hint": "请检查目标文件是否被 Excel 占用、磁盘空间与目录权限"}}, 0)
    except OSError as err:
        emit({"ok": False, "error": {"code": "E1062", "message": "目录表写盘失败: %s" % err, "hint": "请检查目标文件是否被 Excel 占用、磁盘空间与目录权限"}}, 0)
    except Exception as err:  # openpyxl 内部异常（非法字符/合并单元格等）
        emit({"ok": False, "error": {"code": "E1062", "message": "目录表写出失败: %s" % err}}, 0)

    emit({"ok": True, "result": result}, 0)


if __name__ == "__main__":
    main(sys.argv[1:])
